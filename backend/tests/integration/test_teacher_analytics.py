"""Аналитика преподавателя (specs/4.5) через HTTP: норматив / факт, профиль обучающегося, разбор занятия,
подбор задания по слабым местам; чужому преподавателю и обучающемуся — недоступно."""

from collections.abc import Callable, Iterator
from typing import Any

import pytest
from fastapi.testclient import TestClient

from tests.integration.test_assessment import fill_from_reference
from tests.integration.test_journal import app_client as journal_app_client  # noqa: F401
from tests.integration.test_journal import as_user


@pytest.fixture(scope="module")
def app_client(journal_app_client: Callable[[], TestClient]) -> Iterator[Callable[[], TestClient]]:  # noqa: F811
    yield journal_app_client


@pytest.fixture(scope="module")
def lesson(app_client: Callable[[], TestClient]) -> dict[str, Any]:
    """Занятие: оператор заполнил карточку по эталону, ДДС приняла системную карточку; всё оценено."""
    teacher, student, petrov = (as_user(app_client, u) for u in ("teacher", "student", "petrov"))
    ids = teacher.post("/api/v1/training/scenarios/generate", json={"count": 2, "groups": [1], "difficulty": 3}).json()[
        "ids"
    ]
    for sid in ids:
        teacher.post(f"/api/v1/training/scenarios/{sid}/approve")
    dds_code = teacher.get(f"/api/v1/training/scenarios/{ids[0]}").json()["reference_card"]["services"][0]["code"]
    people = {s["login"]: s["id"] for s in teacher.get("/api/v1/training/students").json()}
    body = {
        "title": "Аналитика: пожары",
        "mode": "mixed",
        "card_source": "generated",
        "groups": [1],
        "participants": [
            {"student_id": people["student"], "role": "112"},
            {"student_id": people["petrov"], "role": "dds", "dds_service_code": dds_code},
        ],
        "settings": {"feed_interval_s": 60, "max_waiting": 2, "difficulty": 3},
    }
    session_id = teacher.post("/api/v1/training/sessions", json=body).json()["id"]
    teacher.post(f"/api/v1/training/sessions/{session_id}/start")

    fed = petrov.post("/api/v1/training/sessions/feed").json()
    base = f"/api/v1/incidents/dds/{dds_code}/cards/{fed['card_id']}"
    petrov.post(f"{base}/received")
    petrov.post(f"{base}/status", json={"status": "accepted", "order_no": "7", "comment": "Выехали"})

    call = student.post("/api/v1/training/calls/incoming", json={"groups": [1], "difficulty": 3}).json()
    ref = teacher.get(f"/api/v1/training/scenarios/{call['scenario_id']}").json()["reference_card"]
    card = student.post(
        "/api/v1/incidents/cards",
        json={"aon": call["aon"], "scenario_id": call["scenario_id"], "session_id": session_id},
    ).json()
    services = [{"code": s["code"], "is_main": s["main"], "added_by": "auto"} for s in ref["services"]]
    r = student.post(
        f"/api/v1/incidents/cards/{card['id']}/save",
        json={"data": fill_from_reference(ref, call["aon"]), "services": services},
    )
    assert r.status_code == 200, r.text
    assert teacher.post(f"/api/v1/assessment/sessions/{session_id}/evaluate").json()["assessed"] >= 2
    return {"session_id": session_id, "people": people, "card_number": card["number"], "dds": dds_code}


def test_norm_report_compares_fact_with_pp1931(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    teacher = as_user(app_client, "teacher")
    r = teacher.get("/api/v1/assessment/analytics/norms", params={"session_id": lesson["session_id"]})
    assert r.status_code == 200, r.text
    v = r.json()
    assert "1931" in v["source"]
    assert v["card_112"]["count"] == 1 and v["card_112"]["norm_s"] == 75  # норматив по умолчанию — ПП № 1931
    assert v["card_112"]["within"] == 1  # в тесте карточка сохраняется за доли секунды
    assert v["dds"]["count"] >= 1 and v["dds"]["norm_s"] == 30 and v["dds"]["median_s"] is not None
    assert {row["role"] for row in v["by_student"]} == {"112", "dds"}
    assert v["by_session"][0]["session_id"] == lesson["session_id"]
    period = teacher.get("/api/v1/assessment/analytics/norms", params={"days": 7}).json()
    assert period["card_112"]["count"] >= 1 and period["period"] == "последние 7 дн."
    assert teacher.get("/api/v1/assessment/analytics/norms", params={"days": 0}).json()["period"] == "все занятия"


def test_student_profile_across_sessions(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    teacher = as_user(app_client, "teacher")
    r = teacher.get(f"/api/v1/assessment/analytics/students/{lesson['people']['student']}")
    assert r.status_code == 200, r.text
    p = r.json()
    assert p["full_name"] == "Обучающийся Тестовый" and p["roles"] == ["112"] and p["cards"] >= 1
    assert p["points"] and p["points"][-1]["card_number"] == lesson["card_number"]
    assert p["criteria"] and all(c["group"] is not None for c in p["criteria"])
    assert p["coverage"][0]["group_id"] == 1 and p["coverage"][0]["title"]  # группа сценария, а не выбор оператора
    assert all(g["group_id"] != 1 for g in p["not_practiced"])
    assert p["by_difficulty"] and p["by_difficulty"][0]["difficulty"] == 3  # сценарий подобран по сложности занятия
    assert p["card_112"]["count"] >= 1
    other = teacher.get("/api/v1/assessment/analytics/students/00000000-0000-0000-0000-000000000001")
    assert other.status_code == 404


def test_debrief_lists_characteristic_shortcomings(
    app_client: Callable[[], TestClient], lesson: dict[str, Any]
) -> None:
    teacher = as_user(app_client, "teacher")
    r = teacher.get(f"/api/v1/assessment/sessions/{lesson['session_id']}/debrief")
    assert r.status_code == 200, r.text
    d = r.json()
    # карточка оператора тоже уходит в ДДС Петрова: 112 — 1, ДДС — 2 (системная и от оператора);
    # карточку оператора Петров не открывал — она не оценена
    assert d["title"] == "Аналитика: пожары" and d["cards"] == 3 and d["assessed"] == 2
    assert d["best"] and {b["role"] for b in d["best"]} == {"112", "dds"}
    assert all(set(e) >= {"text", "count", "cards", "students"} for e in d["top_errors"])
    assert all(c["advice"] for c in d["weak_criteria"])
    assert d["card_112"]["count"] == 1 and d["dds"]["count"] >= 1


def test_suggest_assignment_by_weak_spots(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    teacher = as_user(app_client, "teacher")
    ids = [lesson["people"]["student"], lesson["people"]["petrov"]]
    r = teacher.get("/api/v1/assessment/analytics/suggest", params={"student_id": ids})
    assert r.status_code == 200, r.text
    s = r.json()
    assert s["mode"] == "mixed" and 1 <= len(s["groups"]) <= 3 and s["title"].startswith("Работа над ошибками")
    assert 1 <= s["difficulty"] <= 5 and s["difficulty_reason"]
    assert {p["role"] for p in s["participants"]} == {"112", "dds"}
    dds = next(p for p in s["participants"] if p["role"] == "dds")
    assert dds["service_code"] == lesson["dds"]
    one = teacher.get("/api/v1/assessment/analytics/suggest", params={"student_id": [ids[0]]}).json()
    assert one["mode"] == "cards_112" and one["title"].endswith("Обучающийся Тестовый")
    assert teacher.get("/api/v1/assessment/analytics/suggest").status_code == 422


def test_readiness_scale_and_protocol_data(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    teacher = as_user(app_client, "teacher")
    r = teacher.get("/api/v1/assessment/analytics/readiness")
    assert r.status_code == 200, r.text
    v = r.json()
    assert v["teacher"] and len(v["scale"]) == 4 and "31.10.2023" in v["source"]
    rows = {(row["full_name"], row["role"]): row for row in v["rows"]}
    op = rows[("Обучающийся Тестовый", "112")]["readiness"]
    assert op["grade"] is None and op["grade_label"] == "недостаточно данных"  # одна карточка из пяти нужных
    one = teacher.get(
        "/api/v1/assessment/analytics/readiness",
        params={"student_id": [lesson["people"]["student"]], "min_cards": 1, "last": 5},
    ).json()
    assert [row["role"] for row in one["rows"]] == ["112"]
    graded = one["rows"][0]["readiness"]
    assert graded["grade"] in (3, 4, 5) and graded["ready"] and graded["status"] == "готов к зачёту"
    assert graded["timing"] == "в срок" and one["rows"][0]["groups"] == 1
    profile = teacher.get(f"/api/v1/assessment/analytics/students/{lesson['people']['student']}").json()
    assert profile["readiness"][0]["readiness"]["grade_label"] == "недостаточно данных"


def test_validation_report(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    teacher = as_user(app_client, "teacher")
    report = teacher.get(f"/api/v1/assessment/sessions/{lesson['session_id']}/report").json()
    op = next(s for s in report["students"] if s["role"] == "112")["cards"][0]
    r = teacher.post(f"/api/v1/assessment/{op['assessment_id']}/override", json={"score": 60, "comment": "Проверка"})
    assert r.status_code == 200, r.text
    v = teacher.get("/api/v1/assessment/analytics/validation").json()
    b = v["benchmark"]
    assert b["scenarios"] >= 2 and b["cases"] > 20
    assert b["detection"] == 1 and b["critical_caught"] == 1 and b["specificity_112"] == 1
    routing = next(g for g in v["generator"] if g["key"] == "routing")
    assert routing["checked"] >= 2 and routing["share"] == 1  # эталон = автоподбор по текущему классификатору
    assert v["expert"]["pairs"] >= 1 and v["expert"]["mae"] is not None
    assert any("ИИ-судья" in n for n in v["notes"])


def test_analytics_is_teacher_only(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    student = as_user(app_client, "student")
    for url in (
        "/api/v1/assessment/analytics/norms",
        f"/api/v1/assessment/analytics/students/{lesson['people']['student']}",
        f"/api/v1/assessment/sessions/{lesson['session_id']}/debrief",
        "/api/v1/assessment/analytics/readiness",
        "/api/v1/assessment/analytics/validation",
    ):
        assert student.get(url).status_code == 403


def test_speech_input_endpoint(app_client: Callable[[], TestClient]) -> None:
    """Голосовой ввод (п. 1.4): без модели речи — выключен и честно отказывает; с моделью — текст реплики."""
    from aiskra.ai.ports import Transcript

    student = as_user(app_client, "student")
    assert student.get("/api/v1/training/speech").json() == {"enabled": False, "tts": False}
    r = student.post("/api/v1/training/speech", files={"audio": ("a.webm", b"x", "audio/webm")})
    assert r.status_code == 422 and r.json()["error"] == "stt_disabled"

    class Whisper:
        enabled = True

        async def transcribe(self, audio: bytes, *, lang: str = "ru", mime: str = "audio/webm") -> Transcript:
            assert mime.startswith("audio/webm")
            return Transcript(text="Что у вас случилось?")

    app = student.app
    services = app.state.services  # type: ignore[attr-defined]
    real, services.stt = services.stt, Whisper()
    try:
        assert student.get("/api/v1/training/speech").json() == {"enabled": True, "tts": False}
        r = student.post("/api/v1/training/speech", files={"audio": ("a.webm", b"\x1aE", "audio/webm;codecs=opus")})
        assert r.status_code == 200 and r.json() == {"text": "Что у вас случилось?"}
        teacherless = TestClient(app)  # без входа — 401
        assert (
            teacherless.post("/api/v1/training/speech", files={"audio": ("a.webm", b"x", "audio/webm")}).status_code
            == 401
        )
    finally:
        services.stt = real


def test_report_xlsx_is_native_excel(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    """П. 10.9: отчёт по занятию — настоящий XLSX, числа числами, шапка на месте."""
    import io

    from openpyxl import load_workbook

    teacher = as_user(app_client, "teacher")
    r = teacher.get(f"/api/v1/assessment/sessions/{lesson['session_id']}/report.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Занятие" and rows[2][:2] == ("ФИО", "Роль")
    card_row = next(row for row in rows[3:] if row[3] == lesson["card_number"])
    assert isinstance(card_row[3], int) and isinstance(card_row[6], (int, float))
    assert (
        as_user(app_client, "student")
        .get(f"/api/v1/assessment/sessions/{lesson['session_id']}/report.xlsx")
        .status_code
        == 403
    )


def test_scenario_edit_runs_grammar_check(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    """П. 3.6: правка сценария преподавателем — проверка грамотности автоматически, замечания в ответе и сценарии."""
    teacher = as_user(app_client, "teacher")
    sid = teacher.get("/api/v1/training/scenarios", params={"status": "approved"}).json()["items"][0]["id"]
    r = teacher.patch(f"/api/v1/training/scenarios/{sid}", json={"what": "Горит горит балкон"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "draft" and body["checked_by"] == "rules"
    assert {g["rule"] for g in body["grammar"]} >= {"повтор слова", "нет знака препинания в конце"}
    saved = teacher.get(f"/api/v1/training/scenarios/{sid}").json()["legend"]["grammar_check"]
    assert saved["by"] == "rules" and len(saved["remarks"]) == len(body["grammar"])
    fixed = teacher.patch(f"/api/v1/training/scenarios/{sid}", json={"what": "Горит балкон."}).json()
    assert fixed["grammar"] == []
    teacher.post(f"/api/v1/training/scenarios/{sid}/approve")


def test_admin_stops_trainer_services(app_client: Callable[[], TestClient], lesson: dict[str, Any]) -> None:
    """П. 2.2: администратор останавливает и запускает поток вызовов 112 и выдачу карточек ДДС."""
    admin, student, petrov = (as_user(app_client, u) for u in ("admin", "student", "petrov"))
    assert admin.get("/api/v1/system/settings").json()["services"] == {"call_stream": 1, "dds_feed": 1}
    r = admin.put("/api/v1/system/settings/services", json={"values": {"call_stream": 0, "dds_feed": 0}})
    assert r.status_code == 200, r.text
    try:
        stopped = student.post("/api/v1/training/calls/incoming", json={"groups": [1]})
        assert stopped.status_code == 422 and stopped.json()["error"] == "service_stopped"
        feed = petrov.post("/api/v1/training/sessions/feed").json()
        assert feed["card_id"] is None and "остановлена администратором" in feed["reason"]
        assert (
            as_user(app_client, "teacher")
            .put("/api/v1/system/settings/services", json={"values": {"call_stream": 1}})
            .status_code
            == 403
        )
    finally:
        admin.put("/api/v1/system/settings/services", json={"values": {"call_stream": 1, "dds_feed": 1}})
    assert student.post("/api/v1/training/calls/incoming", json={"groups": [1]}).status_code == 200

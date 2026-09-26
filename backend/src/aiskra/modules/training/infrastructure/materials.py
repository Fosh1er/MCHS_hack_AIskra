"""Инфраструктура учебных материалов (п. 4.4): SQL-репозиторий, файлы на диске, извлечение текста, выдержки."""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from uuid import UUID
from xml.etree import ElementTree as ET

from sqlalchemy import column, delete, func, select, table
from sqlalchemy.ext.asyncio import AsyncSession

from aiskra.modules.training.application.ports.materials import MaterialRow
from aiskra.modules.training.domain.material import FileType, Material, MaterialKind, best_passages
from aiskra.modules.training.infrastructure.models import MaterialModel
from aiskra.platform.types import as_utc

_users = table("users", column("id"), column("full_name"))


def _entity(r: MaterialModel) -> Material:
    return Material(
        id=r.id,
        title=r.title,
        kind=MaterialKind(r.kind),
        filename=r.filename,
        file_type=FileType(r.file_type),
        size_bytes=r.size_bytes,
        sha256=r.sha256,
        uploaded_by=r.uploaded_by or UUID(int=0),
        text=r.text or "",
        visible=r.visible,
        use_in_prompts=r.use_in_prompts,
        created_at=as_utc(r.created_at),
    )


class SqlMaterialRepository:
    def __init__(self, session: AsyncSession) -> None:
        self._s = session

    async def add(self, m: Material) -> None:
        self._s.add(
            MaterialModel(
                id=m.id,
                title=m.title,
                kind=m.kind.value,
                filename=m.filename,
                file_type=m.file_type.value,
                size_bytes=m.size_bytes,
                sha256=m.sha256,
                text=m.text,
                visible=m.visible,
                use_in_prompts=m.use_in_prompts,
                uploaded_by=m.uploaded_by,
                created_at=m.created_at,
            )
        )
        await self._s.flush()

    async def get(self, material_id: UUID) -> Material | None:
        r = await self._s.get(MaterialModel, material_id)
        return _entity(r) if r else None

    async def by_sha(self, sha256: str) -> Material | None:
        r = (await self._s.execute(select(MaterialModel).where(MaterialModel.sha256 == sha256))).scalars().first()
        return _entity(r) if r else None

    async def save(self, m: Material) -> None:
        r = await self._s.get(MaterialModel, m.id)
        if r is None:
            return
        r.title, r.kind, r.visible, r.use_in_prompts = m.title, m.kind.value, m.visible, m.use_in_prompts
        await self._s.flush()

    async def delete(self, material_id: UUID) -> None:
        await self._s.execute(delete(MaterialModel).where(MaterialModel.id == material_id))
        await self._s.flush()

    async def rows(self, *, q: str, visible_only: bool) -> list[MaterialRow]:
        stmt = (
            select(MaterialModel, func.length(MaterialModel.text), _users.c.full_name)
            .outerjoin(_users, _users.c.id == MaterialModel.uploaded_by)
            .order_by(MaterialModel.kind, MaterialModel.title)
        )
        if visible_only:
            stmt = stmt.where(MaterialModel.visible.is_(True))
        rows = (await self._s.execute(stmt)).all()
        terms = [t for t in re.findall(r"\w{2,}", q.lower())]
        out = []
        for r, chars, name in rows:
            if terms:  # поиск по названию и тексту без учёта регистра (SQLite не приводит кириллицу — фильтр в Python)
                hay = f"{r.title}\n{r.text}".lower()
                if not all(t in hay for t in terms):
                    continue
            out.append(
                MaterialRow(
                    id=r.id,
                    title=r.title,
                    kind=r.kind,
                    filename=r.filename,
                    file_type=r.file_type,
                    size_bytes=r.size_bytes,
                    visible=r.visible,
                    use_in_prompts=r.use_in_prompts,
                    created_at=as_utc(r.created_at),
                    text_chars=int(chars or 0),
                    uploaded_by_name=name,
                )
            )
        return out

    async def prompt_texts(self) -> list[tuple[str, str]]:
        rows = await self._s.execute(
            select(MaterialModel.title, MaterialModel.text).where(MaterialModel.use_in_prompts.is_(True))
        )
        return [(t, x or "") for t, x in rows.all()]


class LocalFileStorage:
    def __init__(self, directory: str | Path) -> None:
        self._dir = Path(directory)

    def put(self, material_id: UUID, data: bytes) -> None:
        self._dir.mkdir(parents=True, exist_ok=True)
        (self._dir / material_id.hex).write_bytes(data)

    def path(self, material_id: UUID) -> Path:
        return self._dir / material_id.hex

    def delete(self, material_id: UUID) -> None:
        (self._dir / material_id.hex).unlink(missing_ok=True)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class DocumentTextExtractor:
    """PDF — pypdf; DOCX — абзацы и таблицы из word/document.xml; XLSX — строки листов через openpyxl."""

    def extract(self, data: bytes, file_type: FileType) -> str:
        if file_type is FileType.PDF:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
        if file_type is FileType.DOCX:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                root = ET.fromstring(z.read("word/document.xml"))
            paras = []
            for p in root.iter(f"{_W}p"):
                text = "".join(t.text or "" for t in p.iter(f"{_W}t")).strip()
                if text:
                    paras.append(text)
            return "\n".join(paras)
        if file_type is FileType.XLSX:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"# {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c not in (None, "")]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        return data.decode("utf-8", errors="replace")


class SqlMaterialContext:
    """Выдержки из материалов «для генерации» по теме происшествия (лексический поиск; эмбеддинги — P2)."""

    def __init__(self, session: AsyncSession) -> None:
        self._s = session

    async def snippets(self, query: str, limit: int = 3) -> list[tuple[str, str]]:
        return best_passages(query, await SqlMaterialRepository(self._s).prompt_texts(), limit)


__all__ = ["DocumentTextExtractor", "LocalFileStorage", "SqlMaterialContext", "SqlMaterialRepository"]

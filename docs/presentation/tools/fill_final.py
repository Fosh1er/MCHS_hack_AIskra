"""Финальная версия презентации АИскры (ЛЦТ-2026): питч + приложение с обязательными слайдами 7–11 шаблона."""
import copy, sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.oxml.ns import qn

SP = Path(sys.argv[1]); SH = SP / "shots2"
prs = Presentation(str(SP / "deck2/work.pptx"))
ALL = list(prs.slides)  # позиции до перестановки: 1–18 исходные, 19–25 — копии-заготовки в порядке создания
S = lambda n: ALL[n - 1]

def shape(slide, sid):
    return next(sh for sh in slide.shapes if sh.shape_id == sid)

def set_paras(sh, items):
    """Абзацы по образцу исходных: i-й новый абзац берёт формат i-го исходного (или последнего). Элемент — строка
    или (жирная метка, обычный текст)."""
    tb = sh.text_frame._txBody
    orig = tb.findall(qn("a:p"))
    tmpl = [copy.deepcopy(p) for p in orig]
    for p in orig:
        tb.remove(p)
    for i, item in enumerate(items):
        p = copy.deepcopy(tmpl[min(i, len(tmpl) - 1)])
        runs = p.findall(qn("a:r"))
        rpr = [copy.deepcopy(r.find(qn("a:rPr"))) for r in runs]
        for ch in list(p):
            if ch.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
                p.remove(ch)
        end = p.find(qn("a:endParaRPr"))
        parts = [(item, None)] if isinstance(item, str) else [(item[0], True), (item[1], False)]
        for k, (text, bold) in enumerate(parts):
            r = copy.deepcopy(runs[0]) if runs else None
            if r is None:
                from lxml import etree
                r = etree.SubElement(p, qn("a:r")); etree.SubElement(r, qn("a:t"))
            for ch in list(r):
                if ch.tag == qn("a:rPr"):
                    r.remove(ch)
            src = rpr[min(k, len(rpr) - 1)] if rpr and rpr[0] is not None else None
            if src is not None:
                new = copy.deepcopy(src)
                if bold is not None:
                    new.set("b", "1" if bold else "0")
                r.insert(0, new)
            r.find(qn("a:t")).text = text
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        tb.append(p)

def remove(sh):
    sh._element.getparent().remove(sh._element)

def picture(slide, path, x, y, w, h):
    iw, ih = Image.open(path).size
    if iw / ih > w / h:
        nh = w * ih / iw; y += (h - nh) / 2; h = nh
    else:
        nw = h * iw / ih; x += (w - nw) / 2; w = nw
    return slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w), Cm(h))

def caption(slide, text, x, y, w):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(0.8))
    tf = tb.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = text; r.font.size = Pt(10); r.font.italic = True
    from pptx.dml.color import RGBColor
    r.font.color.rgb = RGBColor(0x6B, 0x6B, 0x7B)

# ---------- титул питча
s = S(19)
set_paras(shape(s, 3), ["AIskra"])
set_paras(shape(s, 5), ["ИИ-полигон оператора 112 и диспетчера ДДС: ошибаться — на тренажёре, а не на реальном вызове · финал ЛЦТ-2026"])

# ---------- что нового в финале (по образцу «Работа с заявителем», 6 пунктов)
s = S(20)
set_paras(shape(s, 14), ["ЧТО НОВОГО В ФИНАЛЕ"])
items = [
    ("ИИ в реальном времени", "ответ печатается на глазах, голос — с первой фразы; видно модель и скорость"),
    ("Вход одной кнопкой", "эксперт входит в любую роль без пароля; в боевом контуре такого входа нет"),
    ("Модели под задачу", "Qwen 3.8 Flash — реплики, GPT-4o — сценарии и судья; голос и речь с резервом"),
    ("Психология на первом плане", "шкала ECCS в звонке, страница «Методика», источник у каждого критерия"),
    ("Ветвящиеся сценарии", "решения оператора меняют развитие ситуации и её итог"),
    ("Готово к эксплуатации", "CI, e2e по всем ролям, нагрузка на PostgreSQL, режим prod"),
]
heads, descs = [25, 28, 31, 34, 37, 40], [26, 29, 32, 35, 38, 41]
for (h, d), hid, did in zip(items, heads, descs):
    set_paras(shape(s, hid), [h]); set_paras(shape(s, did), [d])
set_paras(shape(s, 42), ["Всё в main и проходит CI. Спецификации — docs/brief/07_Бэклог_финала_техспека.md и specs/F1–F6"])

# ---------- F1 (по образцу «Заявитель в стрессе»)
s = S(21)
set_paras(shape(s, 10), ["ИИ ОТВЕЧАЕТ ВЖИВУЮ"])
set_paras(shape(s, 6), [
    ("Поток ответа. ", "Реплика заявителя печатается по мере генерации, голос звучит с первой фразы, пока модель дописывает остальное."),
    ("Честные метрики. ", "Под ответом — модель, время до первого слова и скорость. Без модели — «ответ по правилам», без имитации."),
    ("Состояние до ответа. ", "Шкала ECCS и причина изменения появляются сразу, как оператор договорил."),
    ("Проверка. ", "Ответ модели сверяется с профилем; не прошёл — плавно заменяется ответом по правилам."),
    ("Перебивание. ", "Оператор заговорил — голос и генерация останавливаются."),
])
sh6 = shape(s, 6); sh6.left, sh6.top, sh6.width, sh6.height = Cm(1.0), Cm(2.9), Cm(18.9), Cm(14.6)
remove(shape(s, 12)); old = shape(s, 11); x, y, w, h = old.left.cm, old.top.cm, old.width.cm, old.height.cm; remove(old)
picture(s, SH / "f1_state.png", x, y, w, h)
caption(s, "снимок со стенда на тестовой модели — время и скорость на демонстрации зависят от модели", x, y + h + 0.1, w)

# ---------- F4 «Методика» (по образцу «Разбор звонка»)
s = S(22)
set_paras(shape(s, 2), ["МЕТОДИКА — В ИНТЕРФЕЙСЕ"])
old = shape(s, 22); x, y, w, h = old.left.cm, old.top.cm, old.width.cm, old.height.cm; remove(old)
picture(s, SH / "f4_crop.png", x, y, w, h + 3.5)
set_paras(shape(s, 23), [
    "Видят преподаватель и обучающийся",
    "15 состояний заявителя: признаки, что делать, что нельзя, как звучит на каждом уровне ECCS.",
    "Научно-методическая основа: ЦЭПП МЧС, профстандарт 12.002, ПП РФ № 1931, ECCS, APCO/NENA — 36 источников со ссылками.",
    "У каждого критерия оценки — «почему это важно» и ссылка на первоисточник.",
    "В отчёте занятия — топ-5 типичных ошибок группы при работе со стрессом.",
])

# ---------- F5 «Ситуация развивается» (по образцу «Эмулятор АРМ»)
s = S(23)
set_paras(shape(s, 3), ["СИТУАЦИЯ РАЗВИВАЕТСЯ"])
set_paras(shape(s, 7), [
    "Ветвящиеся сценарии",
    "Пока едет помощь, ситуация меняется — и итог зависит от решений оператора.",
    "Быстро отправил карточку пожарной охране — пожар потушен в квартире. Опоздал — огонь перекинулся на соседнюю.",
    "Вредный совет «откройте окно» — выход отрезан дымом; верная инструкция «закройте щели мокрой тканью» — люди дождались пожарных.",
    "4 шаблона: пожар, остановка сердца, запах газа, ДТП. Пороги и инструкции — из официальных памяток, у каждого перехода источник.",
    "Итог и путь по развилкам — в разборе звонка и отчёте занятия; вес в оценке задаёт преподаватель.",
])
for pid in (10, 11):
    remove(shape(s, pid))
picture(s, SH / "f5_crop.png", 19.6, 3.0, 12.4, 15.0)

# ---------- F6 «Готово к эксплуатации» (по образцу «Голос и эмоции», 4 карточки)
s = S(24)
plaque = shape(s, 10); plaque.width = Cm(15.0)
t = shape(s, 14); t.width = Cm(14.4); set_paras(t, ["ГОТОВО К ЭКСПЛУАТАЦИИ"])
cards = [
    ("CI на каждое изменение", "Тесты, линтеры, типы; сборка docker с PostgreSQL и nginx с нуля; смоук в Chrome и Firefox."),
    ("e2e по всем ролям", "9 сценариев через интерфейс: оператор, ДДС, преподаватель, администратор, полигон, отказы модели и сети."),
    ("Нагрузка на PostgreSQL", "100 пользователей: p95 — 117 мс при норме 2 с, 218 карточек/с, ошибок 5xx — 0."),
    ("Режим prod", "Не стартует со слабыми паролями и без TLS; демо-вход и Swagger закрыты. Найдено и исправлено 15 ошибок."),
]
for (h, d), hid, did in zip(cards, [29, 32, 35, 38], [30, 33, 36, 39]):
    set_paras(shape(s, hid), [h]); set_paras(shape(s, did), [d])

# ---------- приложение: разделитель
s = S(25)
set_paras(shape(s, 2), ["ПРИЛОЖЕНИЕ"])
set_paras(shape(s, 3), [
    ("Обязательные слайды шаблона ЛЦТ-2026 (7–11): ", "титул, о команде, состав, история, коротко о решении."),
    ("На защите ", "пропускаются: их содержание раскрыто в основной части презентации."),
])

# ---------- обновления существующих слайдов
set_paras(shape(S(5), 3), [
    "Модульный монолит FastAPI + React, PostgreSQL, docker compose; REST API — 117 операций",
    "ИИ за портами: любая OpenAI-совместимая модель (Ollama, vLLM) или офлайн-режим",
    "Эталон карточки и ДДС — правилами из классификатора заказчика, модель пишет только речь",
    "Автооценка по формулам с весами, ИИ-судья, бенчмарк ошибок для проверки достоверности",
    "Заявитель в стрессе: 15 профилей, движок состояния по шкале ECCS, отдельный блок оценки",
    "Голос: распознавание и синтез с эмоцией, ответ ИИ потоком, разговор без рук",
    "514 автотестов и 9 e2e-сценариев в CI, mypy --strict, контроль границ модулей",
])
s = S(12)
set_paras(shape(s, 33), ["Gemini 3.8 Flash TTS: голос роли, эмоция — стилем; резерв — MAI Voice. В контуре — Piper на CPU."])
set_paras(shape(s, 36), ["Микрофон слушает весь звонок, речь — Qwen3-ASR (резерв Whisper), можно перебить, эхо отсекается."])
s = S(15)
set_paras(shape(s, 17), ["p95 — 117 мс"])
set_paras(shape(s, 18), ["100 пользователей на PostgreSQL, данные за год; норма — до 2 с"])
set_paras(shape(s, 19), ["514 тестов + 9 e2e"])
set_paras(shape(s, 20), ["218 карточек в секунду при записи; отчёт по занятию — 4,5 с"])
set_paras(shape(S(16), 29), [
    ("Один docker compose: ", "nginx с TLS, FastAPI, PostgreSQL 16, Ollama — по профилю"),
    ("Сервер по ТЗ: ", "6 ядер, 32 ГБ, без GPU; внешние сети не нужны"),
    ("Модели на демо: ", "реплики — Qwen 3.8 Flash, сценарии и ИИ-судья — GPT-4o; голос — Gemini 3.8 Flash TTS (резерв MAI Voice), речь — Qwen3-ASR (резерв Whisper)"),
    ("В контуре заказчика: ", "любая OpenAI-совместимая модель — Ollama, vLLM, llama.cpp; без модели — офлайн-режим"),
    ("Модель говорит, правила решают: ", "эталон, службы и баллы считают правила; ответ модели — потоком, с проверкой по профилю"),
    ("Режимы prod / demo: ", "в prod — проверка настроек при запуске, без демо-входа; права и аудит на каждом запросе"),
    ("Внедрение: ", "docker compose up → импорт справочников заказчика → учётки и группы → занятие"),
])
s = S(17)
plan = [
    (55, "Пилот", "занятия в учебном центре, калибровка оценки с экспертами (κ, ICC)"),
    (56, "Редактор сценариев", "преподаватель правит граф ветвления; сценарии из обезличенных реальных вызовов"),
    (57, "Дуплексный голос", "оператор и заявитель говорят одновременно, синтез идёт вслед за токенами"),
    (58, "Материалы", "поиск по учебным материалам на эмбеддингах, журнал учёта занятий"),
    (59, "Тиражирование", "ЕДДС районов и другие регионы: меняются только справочники"),
]
for sid, h, d in plan:
    set_paras(shape(s, sid), [h, d])
set_paras(shape(S(18), 3), [
    ("Стенд: ", "адрес — в README; вход в любую роль одной кнопкой"),
    ("Репозиторий: ", "github.com/Fosh1er/MCHS_hack_AIskra — CI на каждое изменение"),
    ("Пояснительная записка: ", "docs/delivery/out — DOCX и PDF"),
    ("Техспека финала: ", "docs/brief/07_Бэклог_финала_техспека.md, specs/F1–F6"),
    ("Руководства пользователя: ", "docs/manuals — преподаватель и обучающийся"),
    ("Исследование и источники: ", "docs/research"),
    ("Проверка: ", "e2e — backend/tests/e2e, чек-лист ручной проверки — docs/qa/bugbash.md"),
])

# ---------- порядок: питч, затем приложение с обязательными слайдами 7–11
order = [19, 6, 7, 20, 21, 9, 10, 22, 23, 11, 12, 8, 13, 14, 24, 15, 16, 17, 18, 25, 1, 2, 3, 4, 5]
lst = prs.slides._sldIdLst
els = list(lst)
for el in els:
    lst.remove(el)
for n in order:
    lst.append(els[n - 1])
prs.save(str(SP / "deck2/final.pptx"))

# PDF: LibreOffice не рисует анимированную гифку — в версии для PDF первый кадр
import io, os
demo = S(7)
gif = next(sh for sh in demo.shapes if sh.shape_type == 13)
frame = Image.open(io.BytesIO(gif.image.blob)).convert("RGB")
poster = SP / "deck2/demo_poster.png"; frame.save(poster)
x, y, w, h = gif.left, gif.top, gif.width, gif.height
remove(gif)
demo.shapes.add_picture(str(poster), x, y, w, h)
prs.save(str(SP / "deck2/final_poster.pptx"))
print("saved", len(prs.slides))
